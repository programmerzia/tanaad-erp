import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Colors
C_DARK = RGBColor(15, 23, 42)        # #0f172a Deep Navy
C_CARD = RGBColor(30, 41, 59)        # #1e293b Slate Card
C_BORDER = RGBColor(51, 65, 85)      # #334155
C_BLUE = RGBColor(14, 165, 233)      # #0ea5e9 Sky Blue
C_GREEN = RGBColor(34, 197, 94)      # #22c55e Emerald
C_AMBER = RGBColor(245, 158, 11)     # #f59e0b Amber
C_WHITE = RGBColor(255, 255, 255)    # White
C_MUTED = RGBColor(148, 163, 184)    # #94a3b8 Slate Muted

def set_slide_bg(slide):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = C_DARK
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, title_text, category_text="TANAAD MANAGEMENT SYSTEM V3"):
    # Category badge
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8), Inches(0.4))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = C_BLUE
    
    # Title
    t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.5), Inches(0.9))
    tf_t = t_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(28)
    p_t.font.bold = True
    p_t.font.color.rgb = C_WHITE

def create_card(slide, left, top, width, height, title, items, border_color=C_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = C_CARD
    card.line.color.rgb = border_color
    card.line.width = Pt(1.5)
    
    txBox = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
    tf = txBox.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_WHITE
    p.space_after = Pt(14)
    
    for item in items:
        p2 = tf.add_paragraph()
        p2.text = "• " + item
        p2.font.size = Pt(13)
        p2.font.color.rgb = C_MUTED
        p2.space_after = Pt(8)

# ==========================================
# SLIDE 1: Title Slide
# ==========================================
s1 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s1)

# Decorative accent bar
bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.8), Inches(1.2), Inches(0.08))
bar.fill.solid()
bar.fill.fore_color.rgb = C_BLUE
bar.line.fill.background()

t_box = s1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(2.5))
tf = t_box.text_frame
tf.word_wrap = True

p1 = tf.paragraphs[0]
p1.text = "TANAAD ERP (v3.0)"
p1.font.size = Pt(44)
p1.font.bold = True
p1.font.color.rgb = C_WHITE
p1.space_after = Pt(8)

p2 = tf.add_paragraph()
p2.text = "Simplified Factory Operations & Real-Time Reporting"
p2.font.size = Pt(24)
p2.font.color.rgb = C_BLUE
p2.space_after = Pt(16)

p3 = tf.add_paragraph()
p3.text = "Built for Kampala Operations: Woven Mats • Tanaad Packages Ltd • HDPE Recycling"
p3.font.size = Pt(15)
p3.font.color.rgb = C_MUTED

# Presenter footer
f_box = s1.shapes.add_textbox(Inches(0.8), Inches(5.8), Inches(10), Inches(0.8))
tf_f = f_box.text_frame
p_f = tf_f.paragraphs[0]
p_f.text = "DEMO & SYSTEM WALKTHROUGH | SINGLE-TAP FACTORY REPORTING"
p_f.font.size = Pt(12)
p_f.font.bold = True
p_f.font.color.rgb = C_GREEN

# ==========================================
# SLIDE 2: Ground Reality & Challenge
# ==========================================
s2 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s2)
add_header(s2, "The Reality On The Factory Floor")

create_card(s2, Inches(0.8), Inches(2.0), Inches(3.6), Inches(4.7), 
    "1. Non-Technical Staff", [
        "Supervisors & cashiers are not software specialists.",
        "They use shared rugged tablets on dusty floor.",
        "Complex ERPs with 50+ menus lead to user confusion & rejection.",
        "Zero training barrier is mandatory."
    ], C_AMBER)

create_card(s2, Inches(4.8), Inches(2.0), Inches(3.6), Inches(4.7), 
    "2. Paper Ledger Delays", [
        "Production & scrap counts kept in exercise books.",
        "Casual daily wages totaled manually with calculators.",
        "End-of-day losses discovered days or weeks later.",
        "Material leakage at dispatch gate unnoticed."
    ], C_BORDER)

create_card(s2, Inches(8.8), Inches(2.0), Inches(3.6), Inches(4.7), 
    "3. The v3 Mandate", [
        "Strip away all unnecessary clutter.",
        "Cut down to 5 essential destinations.",
        "Instant live reports accessible in 1 tap.",
        "CEO oversight from anywhere in Kampala or abroad."
    ], C_GREEN)

# ==========================================
# SLIDE 3: 3 Sites, 1 Connected Group
# ==========================================
s3 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s3)
add_header(s3, "Unified Operations Across 3 Industrial Sites")

create_card(s3, Inches(0.8), Inches(2.0), Inches(3.6), Inches(4.7), 
    "Site A: Woven Plastic Mats", [
        "15 Extrusion Looms running 24/7.",
        "Roll production, pattern shifts & scrap tracking.",
        "Target vs actual output in linear meters.",
        "Live downtime monitoring (power, yarn break)."
    ], C_BLUE)

create_card(s3, Inches(4.8), Inches(2.0), Inches(3.6), Inches(4.7), 
    "Site B: Tanaad Packages", [
        "Woven polypropylene sacks & packaging bags.",
        "Bale counts & conversion efficiency.",
        "Export dispatches to Rwanda & South Sudan.",
        "Strict bundle-to-weight dispatch verification."
    ], C_BLUE)

create_card(s3, Inches(8.8), Inches(2.0), Inches(3.6), Inches(4.7), 
    "Site C: HDPE Recycling", [
        "Incoming scrap plastics intake & washing.",
        "Extruder pelletization & reprocessing.",
        "Internal resin supply back to Sites A & B.",
        "Direct scrap-to-pellet yield ratio."
    ], C_BLUE)

# ==========================================
# SLIDE 4: The 5 Core Hubs of v3
# ==========================================
s4 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s4)
add_header(s4, "The 5 Core Modules: Everything in 1 Tap")

card_w = Inches(2.2)
card_h = Inches(4.7)
gap = Inches(0.2)

modules = [
    ("1. Today's Actions", [
        "Shift checklist",
        "Loom readings",
        "Attendance check",
        "Closing petty cash",
        "Green ready badge"
    ]),
    ("2. Production", [
        "Live 15-Loom grid",
        "Output vs Target",
        "Scrap % (1.8%)",
        "Resin in vs Goods",
        "Downtime reasons"
    ]),
    ("3. Attendance", [
        "Headcount summary",
        "31 In / 2 Absent",
        "Casual wage rates",
        "UGX 288,000 payroll",
        "No ghost workers"
    ]),
    ("4. Cash & Expenses", [
        "Drawer balance",
        "UGX 22.99M live",
        "Spares & food logs",
        "Money In / Out split",
        "Prevents drawer theft"
    ]),
    ("5. Dispatch & Gate", [
        "Lorry loads",
        "Scale vs invoice",
        "-0.2% variance check",
        "Hold on mismatch",
        "Zero gate leakage"
    ])
]

for idx, (m_title, m_items) in enumerate(modules):
    left = Inches(0.8) + idx * (card_w + gap)
    create_card(s4, left, Inches(2.0), card_w, card_h, m_title, m_items, C_BLUE if idx==0 else C_BORDER)

# ==========================================
# SLIDE 5: Leadership Cockpit & Next Step
# ==========================================
s5 = prs.slides.add_slide(prs.slide_layouts[6])
set_slide_bg(s5)
add_header(s5, "Executive Oversight & Live Demonstration")

create_card(s5, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.7), 
    "CEO Multi-Site Cockpit", [
        "Consolidated group health across all 3 sites simultaneously.",
        "Today's Aggregate Revenue: UGX 260.6M across operations.",
        "One-click Morning Brief with anomaly detection.",
        "Mobile-first design: Owner can view from phone anywhere.",
        "Full audit trails with zero tampering possible."
    ], C_GREEN)

create_card(s5, Inches(6.8), Inches(2.0), Inches(5.6), Inches(4.7), 
    "Live Walkthrough Agenda", [
        "Step 1: Shift Supervisor Flow (Today's Actions checklist)",
        "Step 2: Floor Production & Scrap Control (15 Looms)",
        "Step 3: Automated Daily Wages (Zero Ghost Workers)",
        "Step 4: Strict Cash-In-Hand & Petty Cash Ledger",
        "Step 5: Gate Security & Weighbridge Dispatch Clearance",
        "Step 6: CEO Group View & Anomaly Approvals"
    ], C_BLUE)

prs.save("TANAAD_ERP_V3_PRESENTATION.pptx")
print("Successfully generated TANAAD_ERP_V3_PRESENTATION.pptx")
